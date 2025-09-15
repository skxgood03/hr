import datetime
import re
import json
from io import BytesIO

from django.core.paginator import Paginator
from django.shortcuts import render
from django.http import HttpResponse
from rest_framework.response import Response
from rest_framework.views import APIView
import openpyxl
from openpyxl.styles import Font, Alignment

from department.models import Department
from hrms.ResultPageVo import ResultPageVo
from hrms.ResultVo import ResultVo
from personal.models import Personal
from personal.serializer import PersonalSerializer
from station.models import Station
from hr.decorators import admin_required


# Create your views here.
class DateUtils:
    def calculate_age(birthday):
        today = datetime.date.today()
        age = today.year - birthday.year
        if (today.month, today.day) < (birthday.month, birthday.day):
            age -= 1
        return age




class PersonalAddView(APIView):
    def post(self,request):
        identity = request.data.get('identity')
        #先校验身份证合法
        reg = r'^[1-9]\d{5}(16|17|18|19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])\d{3}[0-9Xx]$'
        if not re.match(reg,identity):
            return Response(ResultVo.fail('身份证不合法'))
        #从身份证提取出生日期
        # #612123200504022222
        identify_year = int(identity[6:10])
        identify_month = int(identity[10:12])
        identify_day = int(identity[12:14])
        birthday =datetime.date(identify_year,identify_month,identify_day)

        age = DateUtils.calculate_age(birthday)
        name = request.data.get('name')
        gender = request.data.get('gender')
        email = request.data.get('email')
        phone = request.data.get('phone')
        education = request.data.get('education')
        school = request.data.get('school')
        beginDate =datetime.date.today().strftime('%Y-%m-%d')
        stationId = request.data.get('stationId')
        workStatus = request.data.get('workStatus')
        try:
            station = Station.objects.get(id = stationId)
            msg =Personal.objects.create(name=name,gender=gender,phone=phone,workStatus=workStatus,
                                             birthday=birthday,age=age,identity=identity,education=education,
                                             school=school,beginDate=beginDate,email=email,station=station)
            if msg:
                return Response(ResultVo.success('添加成功'))
            return Response(ResultVo.fail('添加失败'))
        except Station.DoesNotExist:
            return Response(ResultVo.fail('添加失败'))

class PersnoalListsView(APIView):
    def get(self, request):
        # 获取请求参数
        # 页数
        pageNum = request.query_params.get('pageNum', 1)
        # 获取每页显示条数
        pageSize = request.query_params.get('pageSize', 10)
        # 员工姓名
        name = request.query_params.get('name')
        # 工作状态
        workStatus = request.query_params.get('workStatus')
        # 部门ID
        departmentId = request.query_params.get('departmentId')
        
        # 转换为整数
        try:
            pageNum = int(pageNum)
            pageSize = int(pageSize)
        except (ValueError, TypeError):
            pageNum = 1
            pageSize = 10
            
        # 获取所有的员工信息
        personals = Personal.objects.all().order_by("id")
        # 通过员工姓名进行模糊筛查
        if name:
            personals = personals.filter(name__icontains=name)
        #通过工作状态筛查
        if workStatus:
            personals =personals.filter(workStatus=workStatus)
        #通过部门ID去筛查
        if departmentId:
            #先通过部门ID去岗位表里查询该部门下所有岗位的ID
            stationIds = Station.objects.all().filter(department=departmentId).values_list('id')
            #在员工表里通过岗位ID去筛查
            personals =personals.filter(station__in=stationIds)
        # 进行分页
        pagination = Paginator(personals, pageSize)
        page = pagination.page(pageNum)
        #对查询出来的数据进行序列化
        personalsSerializer = PersonalSerializer(page,many=True)
        return Response(ResultPageVo.success("查询成功",pagination.count,personalsSerializer.data))


class PersonalGetByIdView(APIView):
    def get(self,request):
        id = request.query_params.get('id')
        personal = Personal.objects.get(id = id)

        serializer = PersonalSerializer(personal)
        return Response(ResultVo.success('查询成功',serializer.data))

class PersonalUpdateView(APIView):
    def put(self,request):
        identity = request.query_params.get('identity')
        #先校验身份证合法
        reg = r'^[1-9]\d{5}(16|17|18|19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])\d{3}[0-9Xx]$'
        if not re.match(reg,identity):
            return Response(ResultVo.fail('身份证不合法'))
        #从身份证提取出生日期
        # #612123200504022222
        identify_year = int(identity[6:10])
        identify_month = int(identity[10:12])
        identify_day = int(identity[12:14])
        birthday =datetime.date(identify_year,identify_month,identify_day)

        age = DateUtils.calculate_age(birthday)
        name = request.query_params.get('name')
        gender = request.query_params.get('gender')
        email = request.query_params.get('email')
        phone = request.query_params.get('phone')
        education = request.query_params.get('education')
        school = request.query_params.get('school')
        beginDate =datetime.date.today().strftime('%Y-%m-%d')
        stationId = request.query_params.get('stationId')
        workStatus = request.query_params.get('workStatus')
        id = request.query_params.get('id')
        try:
            station = Station.objects.get(id = stationId)
            msg =Personal.objects.filter(id = id).update(name=name,gender=gender,phone=phone,workStatus=workStatus,birthday=birthday,age=age,identity=identity,education=education,school=school,beginDate=beginDate,email=email,station=station)
            if msg:
                return Response(ResultVo.success('更新成功'))
            return Response(ResultVo.fail('更新失败'))
        except Station.DoesNotExist:
            return Response(ResultVo.fail('更新失败'))

class PersonalDeleteView(APIView):
    def delete(self, request):
        id = request.query_params['id']
        try:
            personal = Personal.objects.get(id = id)
            personal.delete()
            return Response(ResultVo.success('删除成功'))
        except Personal.DoesNotExist:
            return Response(ResultVo.fail('删除失败'))

class PersonalNameByDepartmentView(APIView):
    def get(self,request):
        departmentId = request.query_params.get('departmentId')
        personals = Personal.objects.all()
        if departmentId:
            # 先通过部门ID去岗位表里查询该部门下所有岗位的ID
            stationIds = Station.objects.all().filter(department=departmentId).values_list('id')
            # 在员工表里通过岗位ID去筛查
            personals = personals.filter(station__in=stationIds)
        personalsSerializer = PersonalSerializer(personals,many=True)
        return Response(ResultVo.success("查询",personalsSerializer.data))


# 新的个人信息管理视图类（适配前端）
class PersonalWebAddView(APIView):
    def post(self, request):
        try:
            # 使用request.data获取解析后的数据
            data = request.data
            
            # 验证必填字段
            required_fields = ['personalNumber', 'personalName', 'personalSex', 'personalAge', 'personalPhone', 'personalEntryTime', 'personalStatus']
            for field in required_fields:
                if not data.get(field):
                    return Response(ResultVo.fail(f'{field} 不能为空'))
            
            # 检查工号是否已存在
            if Personal.objects.filter(identity=data['personalNumber']).exists():
                return Response(ResultVo.fail('工号已存在'))
            
            # 获取部门和岗位
            department = None
            station = None
            if data.get('departmentId'):
                try:
                    department = Department.objects.get(id=data['departmentId'])
                except Department.DoesNotExist:
                    return Response(ResultVo.fail('部门不存在'))
            
            if data.get('stationId'):
                try:
                    station = Station.objects.get(id=data['stationId'])
                except Station.DoesNotExist:
                    return Response(ResultVo.fail('岗位不存在'))
            
            # 创建员工记录（映射前端字段名到模型字段名）
            personal = Personal.objects.create(
                name=data['personalName'],  # personalName -> name
                gender=1 if data['personalSex'] == '男' else 2,  # personalSex -> gender (1:男，2:女)
                age=data['personalAge'],  # personalAge -> age
                phone=data['personalPhone'],  # personalPhone -> phone
                email=data.get('personalEmail', ''),  # personalEmail -> email
                station=station,
                beginDate=data['personalEntryTime'],  # personalEntryTime -> beginDate
                workStatus=1 if data['personalStatus'] == '在职' else (2 if data['personalStatus'] == '试用期' else 3),  # personalStatus -> workStatus
                # 临时设置必填字段的默认值
                birthday='1990-01-01',  # 默认生日
                identity=data['personalNumber'],  # 使用工号作为临时身份证
                education=1,  # 默认学历
                school='未填写'  # 默认学校
            )
            
            return Response(ResultVo.success('员工添加成功'))
            
        except Exception as e:
            return Response(ResultVo.fail(f'添加失败: {str(e)}'))

class PersonalWebDetailView(APIView):
    def get(self, request, personal_id):
        try:
            personal = Personal.objects.get(id=personal_id)
            
            # 构建返回数据（映射模型字段名到前端期望的字段名）
            personal_data = {
                'id': personal.id,
                'personalNumber': personal.identity,  # identity -> personalNumber
                'personalName': personal.name,  # name -> personalName
                'personalSex': '男' if personal.gender == 1 else '女',  # gender -> personalSex
                'personalAge': personal.age,  # age -> personalAge
                'personalPhone': personal.phone,  # phone -> personalPhone
                'personalEmail': personal.email,  # email -> personalEmail
                'departmentId': personal.station.department.id if personal.station and personal.station.department else None,
                'departmentName': personal.station.department.departmentName if personal.station and personal.station.department else '',
                'stationId': personal.station.id if personal.station else None,
                'stationName': personal.station.stationName if personal.station else '',
                'personalEntryTime': personal.beginDate,  # beginDate -> personalEntryTime
                'personalStatus': '在职' if personal.workStatus == 1 else ('试用期' if personal.workStatus == 2 else '离职'),  # workStatus -> personalStatus
                'personalAddress': ''  # 模型中没有地址字段，返回空字符串
            }
            
            return Response(ResultVo.success('查询成功', personal_data))
        except Personal.DoesNotExist:
            return Response(ResultVo.fail('员工不存在'))
        except Exception as e:
            return Response(ResultVo.fail(f'查询失败: {str(e)}'))

class PersonalWebUpdateView(APIView):
    def post(self, request):
        try:
            # 使用request.data获取解析后的数据
            data = request.data
            
            # 验证必填字段
            if not data.get('id'):
                return Response(ResultVo.fail('员工ID不能为空'))
            
            # 获取员工记录
            try:
                personal = Personal.objects.get(id=data['id'])
            except Personal.DoesNotExist:
                return Response(ResultVo.fail('员工不存在'))
            
            # 检查身份证是否被其他员工使用（使用identity字段）
            if data.get('personalNumber') and Personal.objects.filter(
                identity=data['personalNumber']
            ).exclude(id=data['id']).exists():
                return Response(ResultVo.fail('工号已被其他员工使用'))
            
            # 获取部门和岗位
            department = None
            station = None
            if data.get('departmentId'):
                try:
                    department = Department.objects.get(id=data['departmentId'])
                except Department.DoesNotExist:
                    return Response(ResultVo.fail('部门不存在'))
            
            if data.get('stationId'):
                try:
                    station = Station.objects.get(id=data['stationId'])
                except Station.DoesNotExist:
                    return Response(ResultVo.fail('岗位不存在'))
            
            # 更新员工信息（映射前端字段名到模型字段名）
            update_fields = {}
            if data.get('personalNumber'):
                update_fields['identity'] = data['personalNumber']  # personalNumber -> identity
            if data.get('personalName'):
                update_fields['name'] = data['personalName']  # personalName -> name
            if data.get('personalSex'):
                update_fields['gender'] = 1 if data['personalSex'] == '男' else 2  # personalSex -> gender
            if data.get('personalAge'):
                update_fields['age'] = data['personalAge']  # personalAge -> age
            if data.get('personalPhone'):
                update_fields['phone'] = data['personalPhone']  # personalPhone -> phone
            if 'personalEmail' in data:
                update_fields['email'] = data['personalEmail']  # personalEmail -> email
            if station is not None:
                update_fields['station'] = station
            if data.get('personalEntryTime'):
                update_fields['beginDate'] = data['personalEntryTime']  # personalEntryTime -> beginDate
            if data.get('personalStatus'):
                update_fields['workStatus'] = 1 if data['personalStatus'] == '在职' else (2 if data['personalStatus'] == '试用期' else 3)  # personalStatus -> workStatus
            
            # 执行更新
            for field, value in update_fields.items():
                setattr(personal, field, value)
            personal.save()
            
            return Response(ResultVo.success('员工信息更新成功'))
            
        except Exception as e:
            return Response(ResultVo.fail(f'更新失败: {str(e)}'))

class PersonalWebDeleteView(APIView):
    def post(self, request):
        try:
            # 使用request.data获取解析后的数据
            data = request.data
            
            # 验证员工ID
            if not data.get('id'):
                return Response(ResultVo.fail('员工ID不能为空'))
            
            # 获取并删除员工记录
            try:
                personal = Personal.objects.get(id=data['id'])
                personal.delete()
                return Response(ResultVo.success('员工删除成功'))
            except Personal.DoesNotExist:
                return Response(ResultVo.fail('员工不存在'))
            
        except Exception as e:
            return Response(ResultVo.fail(f'删除失败: {str(e)}'))

# 前端页面视图函数
@admin_required
def personal_list(request):
    """
    个人信息管理页面
    """
    personals = Personal.objects.all().order_by('id')
    departments = Department.objects.all().order_by('id')
    return render(request, 'personal/list.html', {
        'personals': personals,
        'departments': departments
    })

def export_personal_excel(request):
    """
    导出员工信息Excel
    """
    # 获取筛选参数
    name = request.GET.get('name')
    workStatus = request.GET.get('workStatus')
    departmentId = request.GET.get('departmentId')
    
    # 获取员工数据
    personals = Personal.objects.all().order_by('id')
    
    # 应用筛选条件
    if name:
        personals = personals.filter(name__icontains=name)
    if workStatus:
        personals = personals.filter(workStatus=workStatus)
    if departmentId:
        stationIds = Station.objects.filter(department=departmentId).values_list('id', flat=True)
        personals = personals.filter(station__in=stationIds)
    
    # 创建Excel工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "员工信息"
    
    # 设置表头
    headers = ['工号', '姓名', '性别', '年龄', '手机号', '邮箱', '身份证', '学历', '学校', '部门', '岗位', '入职日期', '工作状态']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # 填充数据
    for row, personal in enumerate(personals, 2):
        ws.cell(row=row, column=1, value=personal.id)
        ws.cell(row=row, column=2, value=personal.name)
        ws.cell(row=row, column=3, value='男' if personal.gender == 1 else '女')
        ws.cell(row=row, column=4, value=personal.age)
        ws.cell(row=row, column=5, value=personal.phone)
        ws.cell(row=row, column=6, value=personal.email)
        ws.cell(row=row, column=7, value=personal.identity)
        ws.cell(row=row, column=8, value=personal.education)
        ws.cell(row=row, column=9, value=personal.school)
        ws.cell(row=row, column=10, value=personal.station.department.name if personal.station and personal.station.department else '')
        ws.cell(row=row, column=11, value=personal.station.name if personal.station else '')
        ws.cell(row=row, column=12, value=personal.beginDate.strftime('%Y-%m-%d') if personal.beginDate else '')
        ws.cell(row=row, column=13, value='在职' if personal.workStatus == 1 else '离职')
    
    # 调整列宽
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # 保存到内存
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # 创建HTTP响应
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="员工信息_{datetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    
    return response